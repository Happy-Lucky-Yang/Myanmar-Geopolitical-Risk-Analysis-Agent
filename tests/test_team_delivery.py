"""待审入库契约：原件不变、不能转正、重复运行不重复、有效观测与零值分离。"""
import io
import json
import sqlite3
from pathlib import Path

import numpy as np
import pytest

from scripts.stage_team_delivery import (ReviewStore, asset_for, encoded, extract_document,
                                        fingerprint, geodesic_area, process_documents,
                                        sha256, zone_statistics)


def test_moved_package_checksums_no_overwrite_or_escape(tmp_path):
    from scripts.crosscheck_qc_package import verify_checksums
    folder = tmp_path / '04_复现与校验'
    folder.mkdir()
    asset = tmp_path / '资料.csv'
    asset.write_text('月份,质量\n202001,待审', encoding='utf-8')
    manifest = folder / 'checksums.sha256'
    manifest.write_text(sha256(asset) + '  资料.csv\n', encoding='utf-8')
    assert verify_checksums(tmp_path) == {'files_checked':1, 'mismatches':[]}
    asset.write_text('changed', encoding='utf-8')
    assert verify_checksums(tmp_path)['mismatches'] == ['资料.csv']
    manifest.write_text('a' * 64 + '  ../outside.csv\n', encoding='utf-8')
    with pytest.raises(ValueError, match='越界'):
        verify_checksums(tmp_path)


def test_review_store_idempotent_and_no_promotion(tmp_path):
    source = tmp_path / 'example.md'
    source.write_text('参考材料', encoding='utf-8')
    asset = asset_for(tmp_path, source, sha256(source), 'document')
    store = ReviewStore(tmp_path / 'review')
    try:
        records = [{'kind': 'document_block', 'location': 'line:1', 'text': '参考材料', 'eligible_for_scoring': True}]
        assert store.save(asset, records) == 1
        assert store.save(asset, records) == 0
        assert store.counts() == {'assets': 1, 'records': {'document_block': 1}, 'scoring_eligible': 0}
        payload = json.loads(store.conn.execute('SELECT payload FROM pending_records').fetchone()[0])
        assert payload['eligible_for_scoring'] is False
        with pytest.raises(sqlite3.IntegrityError):
            with store.conn:
                store.conn.execute('UPDATE pending_records SET eligible_for_scoring=1')
        with pytest.raises(sqlite3.IntegrityError):
            with store.conn:
                store.conn.execute("UPDATE pending_records SET status='approved'")
        with pytest.raises(ValueError, match='冲突'):
            store.save({**asset, 'path': 'changed'}, records)
    finally:
        store.close()


def test_docx_paragraph_table_extraction_keeps_original(tmp_path):
    from docx import Document
    path = tmp_path / '材料.docx'
    doc = Document()
    doc.add_paragraph('方法：不补零。')
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = '缺测'
    table.cell(0, 1).text = '不等于0'
    doc.save(path)
    digest = sha256(path)
    result = extract_document(path, path.read_bytes())
    assert any(b['text'] == '方法：不补零。' for b in result['blocks'])
    assert any(b['text'] == '不等于0' and b['table_cell'] for b in result['blocks'])
    assert sha256(path) == digest


def test_xlsx_formula_not_executed_and_hidden_sheet_preserved(tmp_path):
    module = pytest.importorskip('openpyxl')
    path = tmp_path / '权重.xlsx'
    book = module.Workbook()
    book.active.append(['数据', 0, '=1/0'])
    hidden = book.create_sheet('复核')
    hidden.sheet_state = 'hidden'
    hidden.append(['待核验'])
    book.save(path)
    digest = sha256(path)
    result = extract_document(path, path.read_bytes())
    cells = json.loads(result['blocks'][0]['text'])
    assert cells[1]['value'] == 0
    assert cells[2]['formula'] and cells[2]['cached'] is None
    assert result['metadata']['sheets'][1]['state'] == 'hidden'
    assert sha256(path) == digest


def test_pptx_text_table_and_notes(tmp_path):
    module = pytest.importorskip('pptx')
    from pptx.util import Inches
    path = tmp_path / '汇报.pptx'
    deck = module.Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text = '只是方案，尚未验收'
    slide.notes_slide.notes_text_frame.text = '不可当正式评分依据'
    deck.save(path)
    result = extract_document(path, path.read_bytes())
    assert result['metadata']['slides'] == 1
    assert any(b['text'] == '只是方案，尚未验收' for b in result['blocks'])
    assert any(b['location'] == 'slide:1:notes' for b in result['blocks'])


def test_pixel_statistics_preserve_negative_and_real_zero():
    rad = np.array([[100, 0, -2], [-9999, 4, np.nan]], dtype='float32')
    cf = np.array([[0, 1, 2], [-9999, 4, 1]], dtype='float32')
    zone = np.ones((2, 3), dtype=bool)
    zone[1, 0] = False
    area = np.array([[1, 1, 1], [1, 2, 1]], dtype='float64')
    result = zone_statistics(rad, cf, zone, area, -9999)
    assert result['inside_pixels'] == 5
    assert result['observed_pixels'] == 3  # 真实零值保留，cf=0亮值不入均值
    assert result['zero_coverage_pixels'] == 1
    assert result['invalid_pixels'] == 1
    assert result['negative_pixels'] == 1
    assert result['radiance_mean_observed'] == pytest.approx(2 / 3)
    assert result['radiance_mean_area_weighted'] == pytest.approx(1.5)
    assert result['cf_mean_observed'] == pytest.approx(7 / 3)


def test_geodesic_area_hole_and_orientation():
    from shapely.geometry import Polygon
    from shapely.geometry.polygon import orient
    outer = [(95, 18), (96, 18), (96, 19), (95, 19)]
    hole = [(95.2, 18.2), (95.8, 18.2), (95.8, 18.8), (95.2, 18.8)]
    polygon = Polygon(outer, [hole])
    assert geodesic_area(polygon) == pytest.approx(geodesic_area(orient(polygon, sign=-1)))
    assert geodesic_area(polygon) == pytest.approx(geodesic_area(Polygon(outer)) - geodesic_area(Polygon(hole)))


def test_dry_run_documents_does_not_create_database(tmp_path):
    source = tmp_path / 'input'
    source.mkdir()
    (source / '.env').write_text('DO_NOT_READ=fixture', encoding='utf-8')
    night = source / '夜间灯光'
    night.mkdir()
    text = night / '说明.md'
    text.write_text('真实零值不是缺测', encoding='utf-8')
    output = tmp_path / 'review'
    digest = sha256(text)
    first = process_documents(source, output)
    assert len(first) == 1
    assert first[0]['blocks'] == 1
    assert not (output / 'review.sqlite3').exists()
    assert sha256(text) == digest
    assert first == process_documents(source, output)


def test_qc_package_flags_partial_months_and_checks_hashes(tmp_path):
    """非整月产品必须被标出并隔离；清单哈希与已入库资产逐条比对。"""
    from scripts.stage_team_delivery import qc_package_audit
    package = tmp_path / 'QC'
    folder = package / '02_三项补充' / '01_原始产品文件名清单'
    folder.mkdir(parents=True)
    manifest = folder / 'source_manifest.csv'
    manifest.write_text(
        'ym,source_file,period_days,period_flag,status,tif_sha256\n'
        '202405,SVDNB_npp_20240501-20240528_75N060E_vcmslcfg_v10_c1.tgz,28,partial,delivered,' + 'a' * 64 + '\n'
        '202406,SVDNB_npp_20240604-20240630_75N060E_vcmslcfg_v10_c2.tgz,27,partial,delivered,' + 'b' * 64 + '\n'
        '202407,SVDNB_npp_20240701-20240731_75N060E_vcmslcfg_v10_c3.tgz,31,full,delivered,' + 'c' * 64 + '\n',
        encoding='utf-8')
    output = tmp_path / 'team_delivery_review'
    store = ReviewStore(output)
    try:
        raster = tmp_path / 'MMR_VNL_202405.tif'
        raster.write_bytes(b'fixture')
        asset = {**asset_for(tmp_path, raster, 'a' * 64, 'raster')}
        store.save(asset, [{'kind': 'nightlight_candidate', 'location': '202405:MMR', 'status': 'pending_review',
                            'region': 'MMR', 'period_start': '2024-05-01', 'period_end': '2024-06-01',
                            'statistics': {}, 'flags': []}])
        result = qc_package_audit(tmp_path, output, package, store)
        assert result['partial_period_months'] == ['202405', '202406']
        assert result['tif_hash_matched'] == 1 and result['tif_hash_mismatched'] == []
        assert result['corrections_recorded'] == 2
        assert qc_package_audit(tmp_path, output, package, store)['corrections_recorded'] == 0
        rows = [json.loads(r[0]) for r in store.conn.execute(
            "SELECT payload FROM pending_records WHERE kind='period_correction'")]
        assert {r['location'] for r in rows} == {'202405', '202406'}
        assert all(r['status'] == 'quarantined' and r['eligible_for_scoring'] is False for r in rows)
        june = next(r for r in rows if r['location'] == '202406')
        assert june['actual_product_start'] == '2024-06-04' and june['actual_product_days'] == 27
        may = next(r for r in rows if r['location'] == '202405')
        assert may['previously_recorded_as_full_month'] is True
        assert may['previously_recorded_period'] == ['2024-05-01', '2024-06-01']
        # 天数与文件名不自洽时必须报错，不默认信任清单列。
        manifest.write_text('ym,source_file,period_days,period_flag,status,tif_sha256\n'
                            '202405,SVDNB_npp_20240501-20240528_75N060E_vcmslcfg_v10_c1.tgz,31,partial,delivered,'
                            + 'a' * 64 + '\n', encoding='utf-8')
        with pytest.raises(ValueError, match='天数'):
            qc_package_audit(tmp_path, output, package, store)
    finally:
        store.close()


def test_raster_audit_commit_rejects_changed_input_and_retries_safely(tmp_path):
    rasterio = pytest.importorskip('rasterio')
    from rasterio.transform import from_origin
    from shapely.geometry import box, mapping
    from scripts.stage_team_delivery import process_rasters, commit_raster_audit, export_json, VERSION
    root = tmp_path / 'input'
    package = root / '夜间灯光' / 'fixture'
    main, bbox, boundary = package / 'data_mmr', package / 'data', package / 'boundary'
    for p in (main, bbox, boundary):
        p.mkdir(parents=True)
    geometry = mapping(box(95, 18, 97, 20))
    national = {'type': 'FeatureCollection', 'features': [{'type': 'Feature', 'properties': {}, 'geometry': geometry}]}
    provincial = {'type': 'FeatureCollection', 'features': [{'type': 'Feature', 'properties': {'GID_1': 'MMR.1_1', 'NAME_1': 'fixture'}, 'geometry': geometry}]}
    (boundary / 'gadm41_MMR_0.json').write_text(encoded(national), encoding='utf-8')
    (boundary / 'gadm41_MMR_1.json').write_text(encoded(provincial), encoding='utf-8')
    transform = from_origin(94, 21, 1, 1)
    raw = np.ones((2, 4, 4), dtype='float32')
    clipped = np.full_like(raw, -9999)
    clipped[:, 1:3, 1:3] = raw[:, 1:3, 1:3]
    for directory, array, nodata in ((main, clipped, -9999), (bbox, raw, None)):
        with rasterio.open(directory / 'MMR_VNL_202001.tif', 'w', driver='GTiff', width=4, height=4,
                           count=2, dtype='float32', transform=transform, crs='EPSG:4326', nodata=nodata) as dst:
            dst.write(array)
    output = tmp_path / 'team_delivery_review'
    result = process_rasters(root, output)
    assert result['months'][0]['equal_bbox']
    assert result['months'][0]['national']['observed_pixels'] == 4
    summary = {'version': VERSION, 'phase': 'rasters', 'result': result}
    audit_path = Path(export_json(output, 'rasters-' + fingerprint(summary)[:16] + '.json', summary))
    store = ReviewStore(output)
    try:
        first = commit_raster_audit(root, output, audit_path, store)
        assert first['new_records'] > 0
        assert commit_raster_audit(root, output, audit_path, store)['new_records'] == 0
        assert store.counts()['scoring_eligible'] == 0
        with rasterio.open(main / 'MMR_VNL_202001.tif', 'r+') as dst:
            dst.update_tags(changed='fixture')
        with pytest.raises(ValueError, match='原件已变化'):
            commit_raster_audit(root, output, audit_path, store)
    finally:
        store.close()
