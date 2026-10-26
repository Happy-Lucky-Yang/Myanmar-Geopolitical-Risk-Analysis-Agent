"""团队交付初步入库：原件只读，Office正文与夜光统计进入独立SQLite待审库。

依赖（保持主项目NumPy版本）：
python -m pip install numpy==1.26.4 rasterio==1.4.3 openpyxl==3.1.5 python-pptx==1.0.2
默认只输出核验JSON；--apply才写review.sqlite3。绝不连接正式指标、新闻或RAG库。
"""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import date, datetime, timezone
import hashlib
import io
import json
import os
from pathlib import Path
import re
import sqlite3
import sys
import zipfile

VERSION = 'team-delivery-review-v1'
OFFICE = {'.docx', '.pptx', '.xlsx'}


def encoded(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, allow_nan=False, default=str)


def fingerprint(value):
    return hashlib.sha256(encoded(value).encode('utf-8')).hexdigest()


def sha256(path):
    h = hashlib.sha256()
    with Path(path).open('rb') as f:
        for block in iter(lambda: f.read(1024 * 1024), b''):
            h.update(block)
    return h.hexdigest()


def safe_files(root, suffixes):
    root = Path(root).resolve(strict=True)
    for current, dirs, files in os.walk(root, followlinks=False):
        dirs[:] = [d for d in sorted(dirs) if not d.startswith('.') and
                   not Path(current, d).is_symlink() and not Path(current, d).is_junction()]
        for name in sorted(files):
            path = Path(current, name)
            if name.startswith('~$') or path.is_symlink() or path.suffix.lower() not in suffixes:
                continue
            if path.resolve().is_relative_to(root):
                yield path


def document_paths(root):
    # 只摄取用户指定的交付材料；不扫描运行数据、账户、Cookie或环境文件。
    result = [p for p in root.iterdir() if p.is_file() and not p.is_symlink() and p.suffix.lower() in OFFICE]
    theory = root / '（中期）理论部分材料'
    if theory.exists():
        result.extend(safe_files(theory, OFFICE))
    night = root / '夜间灯光'
    if night.exists():
        result.extend(safe_files(night, {'.md', '.csv'}))
    return sorted(set(result))


def inspect_archive(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if sum(i.file_size for i in archive.infolist()) > 512 * 1024 * 1024:
            raise ValueError('Office解压体积过大，需单独核验')
        return archive.namelist()


def extract_document(path, data):
    """提取正文、表格、备注和公式缓存；不执行宏、不计算公式、不OCR图片。"""
    suffix = path.suffix.lower()
    blocks, meta = [], {'image_ocr': False, 'formula_evaluation': False}
    if suffix in {'.md', '.csv'}:
        text = data.decode('utf-8-sig')
        blocks = [{'location': f'line:{n}', 'text': line} for n, line in enumerate(text.splitlines(), 1) if line.strip()]
    elif suffix == '.docx':
        from lxml import etree
        names = inspect_archive(data)
        ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
              'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math'}
        parser = etree.XMLParser(resolve_entities=False, no_network=True)
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            parts = [n for n in names if n == 'word/document.xml' or
                     re.fullmatch(r'word/(header\d+|footer\d+|footnotes|endnotes|comments)\.xml', n)]
            for part in sorted(parts):
                tree = etree.fromstring(archive.read(part), parser)
                # 表格单元格的段落也按XML顺序读取；OMML公式保留结构，避免扁平化丢运算符。
                for i, para in enumerate(tree.xpath('//w:p', namespaces=ns), 1):
                    text = ''.join(para.xpath('.//w:t/text() | .//m:t/text()', namespaces=ns))
                    maths = para.xpath('.//m:oMath', namespaces=ns)
                    if text.strip() or maths:
                        blocks.append({'location': f'{part}:paragraph:{i}', 'text': text,
                                       'table_cell': bool(para.xpath('ancestor::w:tc', namespaces=ns)),
                                       'math_xml': [etree.tostring(m, encoding='unicode') for m in maths]})
        meta.update(images=sum(n.startswith('word/media/') for n in names),
                    embedded_objects=sum(n.startswith('word/embeddings/') for n in names))
    elif suffix == '.pptx':
        from pptx import Presentation
        names = inspect_archive(data)
        presentation = Presentation(io.BytesIO(data))

        def shapes(items, prefix):
            for i, shape in enumerate(items, 1):
                location = f'{prefix}:shape:{i}'
                if shape.has_text_frame and shape.text.strip():
                    blocks.append({'location': location, 'text': shape.text})
                if shape.has_table:
                    for r, row in enumerate(shape.table.rows, 1):
                        blocks.append({'location': f'{location}:row:{r}',
                                       'text': ' | '.join(c.text for c in row.cells)})
                if shape.has_chart:
                    for s, series in enumerate(shape.chart.series, 1):
                        blocks.append({'location': f'{location}:series:{s}',
                                       'text': encoded({'name': series.name, 'values': list(series.values)})})
                if hasattr(shape, 'shapes'):
                    shapes(shape.shapes, location)

        for i, slide in enumerate(presentation.slides, 1):
            shapes(slide.shapes, f'slide:{i}')
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                text = slide.notes_slide.notes_text_frame.text
                if text.strip():
                    blocks.append({'location': f'slide:{i}:notes', 'text': text})
        meta.update(slides=len(presentation.slides), images=sum(n.startswith('ppt/media/') for n in names),
                    embedded_objects=sum(n.startswith('ppt/embeddings/') for n in names))
    elif suffix == '.xlsx':
        from openpyxl import load_workbook
        inspect_archive(data)
        book = load_workbook(io.BytesIO(data), read_only=False, data_only=False, keep_links=False)
        cache = load_workbook(io.BytesIO(data), read_only=False, data_only=True, keep_links=False)
        sheets = []
        try:
            for sheet in book:
                count = 0
                for row in sheet.iter_rows():
                    cells = []
                    for cell in row:
                        if cell.value is not None:
                            cells.append({'cell': cell.coordinate, 'value': cell.value,
                                          'formula': cell.data_type == 'f',
                                          'cached': cache[sheet.title][cell.coordinate].value if cell.data_type == 'f' else None,
                                          'hyperlink': cell.hyperlink.target if cell.hyperlink else None})
                    if cells:
                        count += 1
                        blocks.append({'location': f'{sheet.title}:row:{row[0].row}', 'text': encoded(cells)})
                sheets.append({'name': sheet.title, 'state': sheet.sheet_state, 'populated_rows': count,
                               'merged_ranges': [str(r) for r in sheet.merged_cells.ranges]})
            meta['sheets'] = sheets
        finally:
            book.close()
            cache.close()
    else:
        raise ValueError('不支持的资料格式')
    return {'metadata': meta, 'blocks': blocks}


class ReviewStore:
    """只含待审表，无正式观测表名，无自动转正入口。"""
    def __init__(self, directory):
        directory = Path(directory)
        if directory.is_symlink() or directory.is_junction():
            raise ValueError('待审目录不能是链接')
        directory.mkdir(parents=True, exist_ok=True)
        self.path = directory / 'review.sqlite3'
        if self.path.is_symlink():
            raise ValueError('待审库不能是链接')
        self.conn = sqlite3.connect(self.path)
        self.conn.execute('PRAGMA foreign_keys=ON')
        self.conn.execute('CREATE TABLE IF NOT EXISTS assets (id TEXT PRIMARY KEY, path TEXT NOT NULL, sha256 TEXT NOT NULL, kind TEXT NOT NULL, payload TEXT NOT NULL, created_at TEXT NOT NULL)')
        self.conn.execute('CREATE TABLE IF NOT EXISTS pending_records (id TEXT PRIMARY KEY, asset_id TEXT NOT NULL REFERENCES assets(id), kind TEXT NOT NULL, location TEXT NOT NULL, status TEXT NOT NULL CHECK(status IN (\'pending_review\',\'quarantined\',\'missing\')), eligible_for_scoring INTEGER NOT NULL CHECK(eligible_for_scoring=0), payload TEXT NOT NULL)')
        self.conn.commit()

    def save(self, asset, records):
        inserted = 0
        with self.conn:
            previous = self.conn.execute('SELECT payload FROM assets WHERE id=?', (asset['id'],)).fetchone()
            payload = encoded(asset)
            if previous and previous[0] != payload:
                raise ValueError('待审资产内容冲突，拒绝覆盖')
            self.conn.execute('INSERT OR IGNORE INTO assets VALUES (?,?,?,?,?,?)',
                              (asset['id'], asset['path'], asset['sha256'], asset['kind'], payload,
                               datetime.now(timezone.utc).isoformat()))
            for record in records:
                record = {**record, 'eligible_for_scoring': False, 'review_version': VERSION}
                if record['kind'] == 'nightlight_candidate':
                    reasons = [f for f in record.get('flags', []) if f.startswith('delivery_declared_anomaly_')
                               or f in {'invalid_pixels', 'coverage_below_50pct_review_only'}]
                    record['quarantine_reasons'] = reasons
                    if reasons:
                        record['status'] = 'quarantined'
                identity = fingerprint([asset['id'], record])
                cursor = self.conn.execute('INSERT OR IGNORE INTO pending_records VALUES (?,?,?,?,?,?,?)',
                                           (identity, asset['id'], record['kind'], record['location'],
                                            record.get('status', 'pending_review'), 0, encoded(record)))
                inserted += cursor.rowcount
        return inserted

    def counts(self):
        return {'assets': self.conn.execute('SELECT count(*) FROM assets').fetchone()[0],
                'records': dict(self.conn.execute('SELECT kind,count(*) FROM pending_records GROUP BY kind')),
                'scoring_eligible': self.conn.execute('SELECT count(*) FROM pending_records WHERE eligible_for_scoring<>0').fetchone()[0]}

    def close(self):
        self.conn.close()


def asset_for(root, path, digest, kind):
    item = {'path': path.relative_to(root).as_posix(), 'sha256': digest, 'kind': kind,
            'bytes': path.stat().st_size, 'review_version': VERSION}
    return {**item, 'id': fingerprint(item)}


def export_json(directory, name, value):
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / name
    text = encoded(value) + '\n'
    if path.exists():
        if path.read_text(encoding='utf-8') != text:
            raise ValueError('不可变核验结果冲突：' + name)
    else:
        with path.open('x', encoding='utf-8') as f:
            f.write(text)
    return str(path)


def process_documents(root, output, store=None):
    results = []
    for path in document_paths(root):
        data = path.read_bytes()
        digest = hashlib.sha256(data).hexdigest()
        result = extract_document(path, data)
        if sha256(path) != digest:
            raise ValueError('原件在读取时变化，请冻结后重试')
        asset = asset_for(root, path, digest, 'document')
        records = [{'kind': 'document_block', 'status': 'pending_review', **block} for block in result['blocks']]
        report = export_json(output / 'documents', asset['id'] + '.json', {'asset': asset, **result})
        transcript = output / 'documents' / (asset['id'] + '.txt')
        readable = '\n'.join(f"[{b['location']}] {b['text']}" for b in result['blocks']) + '\n'
        if not transcript.exists():
            with transcript.open('x', encoding='utf-8') as f:
                f.write(readable)
        elif transcript.read_text(encoding='utf-8') != readable:
            raise ValueError('正文导出发生冲突')
        new = store.save(asset, records) if store else 0
        results.append({'file': asset['path'], 'sha256': digest, 'blocks': len(records),
                        'metadata': result['metadata'], 'report': report, 'new_records': new})
        print(f"正文 {path.name}: {len(records)}块", flush=True)
    return results


def geodesic_area(geom):
    from pyproj import Geod
    from shapely.geometry.polygon import orient
    geod = Geod(ellps='WGS84')
    if geom.is_empty:
        return 0.0
    if geom.geom_type == 'Polygon':
        return abs(geod.geometry_area_perimeter(orient(geom, sign=1))[0])
    return sum(geodesic_area(part) for part in getattr(geom, 'geoms', []))


def zone_statistics(rad, cf, zone, areas, nodata):
    import numpy as np
    finite = np.isfinite(rad) & np.isfinite(cf)
    if nodata is not None:
        finite &= (rad != nodata) & (cf != nodata)
    observed = zone & finite & (cf > 0) & (cf == np.floor(cf))
    n, valid_n = int(zone.sum()), int(observed.sum())
    area = float(areas[zone].sum())
    valid_area = float(areas[observed].sum())
    rv, cv = rad[observed], cf[observed]
    return {'inside_pixels': n, 'observed_pixels': valid_n,
            'zero_coverage_pixels': int((zone & finite & (cf == 0)).sum()),
            'invalid_pixels': int((zone & (~finite | (cf < 0) | (cf != np.floor(cf)))).sum()),
            'observed_pixel_fraction': valid_n / n if n else None,
            'observed_area_fraction': valid_area / area if area else None,
            'negative_pixels': int((rv < 0).sum()),
            'negative_fraction_observed': float((rv < 0).mean()) if valid_n else None,
            'cf_mean_all_inside': float(cf[zone & finite].mean(dtype=np.float64)) if (zone & finite).any() else None,
            'cf_mean_observed': float(cv.mean(dtype=np.float64)) if valid_n else None,
            'radiance_mean_observed': float(rv.mean(dtype=np.float64)) if valid_n else None,
            'radiance_mean_area_weighted': float(np.sum(rv * areas[observed]) / valid_area) if valid_area else None,
            'radiance_min': float(rv.min()) if valid_n else None,
            'radiance_max': float(rv.max()) if valid_n else None,
            'radiance_p50': float(np.median(rv)) if valid_n else None}


def grid_of(src):
    return {'crs': src.crs.to_string() if src.crs else None, 'height': src.height,
            'width': src.width, 'count': src.count, 'transform': list(src.transform),
            'dtypes': list(src.dtypes), 'nodata': src.nodata, 'scales': list(src.scales),
            'offsets': list(src.offsets), 'descriptions': list(src.descriptions), 'units': list(src.units)}


def review_flags(month, stats):
    flags = ['product_provenance_pending', 'aggregation_method_pending']
    if stats['invalid_pixels']:
        flags.append('invalid_pixels')
    if (stats['observed_area_fraction'] or 0) < .5:
        flags.append('coverage_below_50pct_review_only')
    if stats['negative_pixels']:
        flags.append('negative_radiance_preserved')
    if month in {'202001', '202207', '202502', '202511'}:
        flags.append('delivery_declared_anomaly_' + month)
    return flags


def process_rasters(root, output, store=None):
    import numpy as np
    import rasterio
    from rasterio.features import geometry_mask, rasterize
    from shapely.geometry import shape, box, mapping
    from pyproj import Geod
    night = root / '夜间灯光'
    vectors = list(safe_files(night, {'.tif'}))
    main = sorted(p for p in vectors if p.parent.name == 'data_mmr')
    bbox_files = {}
    for p in vectors:
        if p.parent.name == 'data':
            if p.name in bbox_files:
                raise ValueError('同月有多个规则版，需明确版本')
            bbox_files[p.name] = p
    if not main:
        raise ValueError('找不到矢量裁剪影像')
    boundary = main[0].parent.parent / 'boundary'
    boundary_paths = [boundary / f'gadm41_MMR_{i}.json' for i in (0, 1)]
    boundary_hashes = {p.name: sha256(p) for p in boundary_paths}
    national = shape(json.loads(boundary_paths[0].read_text(encoding='utf-8'))['features'][0]['geometry'])
    features = json.loads(boundary_paths[1].read_text(encoding='utf-8'))['features']
    if not national.is_valid or any(not shape(f['geometry']).is_valid for f in features):
        raise ValueError('边界几何无效，禁止静默修复')
    zones = [{'region': 'MMR', 'name': '缅甸全国', 'geometry': national}]
    zones += [{'region': f['properties']['GID_1'], 'name': f['properties']['NAME_1'], 'geometry': shape(f['geometry'])} for f in features]
    with rasterio.open(main[0]) as src:
        reference = grid_of(src)
        if src.crs.to_epsg() != 4326 or src.count != 2 or src.transform.b or src.transform.d or src.transform.a <= 0 or src.transform.e >= 0:
            raise ValueError('只支持WGS84北向上双波段网格')
        inside = geometry_mask([mapping(national)], (src.height, src.width), src.transform, invert=True, all_touched=False)
        labels = rasterize([(mapping(z['geometry']), i) for i, z in enumerate(zones[1:], 1)],
                           out_shape=inside.shape, transform=src.transform, all_touched=False, dtype='uint8')
        geod = Geod(ellps='WGS84')
        weights = []
        for row in range(src.height):
            x, y = src.transform * (0, row)
            dx, dy = src.transform.a, src.transform.e
            a, _ = geod.polygon_area_perimeter([x, x+dx, x+dx, x], [y, y, y+dy, y+dy])
            weights.append(abs(a))
        areas = np.broadcast_to(np.array(weights)[:, None], inside.shape)
        rectangle = box(*src.bounds)
        coverage = [{'region': z['region'], 'name': z['name'],
                     'boundary_area_km2': geodesic_area(z['geometry']) / 1e6,
                     'outside_raster_km2': geodesic_area(z['geometry'].difference(rectangle)) / 1e6}
                    for z in zones]
    boundaries = {'hashes': boundary_hashes, 'method': 'WGS84 ellipsoidal geodesic polygon area; not EPSG:3348 or UTM',
                  'coverage': coverage, 'unassigned_inside_pixels': int((inside & (labels == 0)).sum())}
    results, present = [], set()
    for index, path in enumerate(main, 1):
        match = re.fullmatch(r'MMR_VNL_(\d{6})\.tif', path.name)
        if not match:
            raise ValueError('无法识别月份')
        month = match.group(1)
        if month in present:
            raise ValueError('矢量版月份重复')
        present.add(month)
        digest = sha256(path)
        bbox = bbox_files.get(path.name)
        bbox_digest = sha256(bbox) if bbox else None
        with rasterio.open(path) as src:
            grid = grid_of(src)
            if grid != reference or grid['scales'] != [1.0, 1.0] or grid['offsets'] != [0.0, 0.0]:
                raise ValueError(f'{month}网格或波段语义不一致，停止该批次')
            rad, cf = src.read()
            equal = False
            if bbox:
                with rasterio.open(bbox) as original:
                    original_grid = grid_of(original)
                    for key in ('crs', 'height', 'width', 'count', 'transform', 'scales', 'offsets'):
                        if grid[key] != original_grid[key]:
                            raise ValueError(f'{month}规则版网格不一致')
                    br, bc = original.read()
                    equal = bool(np.array_equal(rad[inside], br[inside], equal_nan=True) and
                                 np.array_equal(cf[inside], bc[inside], equal_nan=True))
                    del br, bc
            outside_correct = bool(np.all(rad[~inside] == src.nodata) and np.all(cf[~inside] == src.nodata))
            records = []
            for i, zone in enumerate(zones):
                stats = zone_statistics(rad, cf, inside if i == 0 else inside & (labels == i), areas, src.nodata)
                start = date(int(month[:4]), int(month[4:]), 1)
                end = date(start.year + (start.month == 12), start.month % 12 + 1, 1)
                records.append({'kind': 'nightlight_candidate', 'location': f'{month}:{zone["region"]}',
                                'status': 'pending_review' if equal and outside_correct else 'quarantined',
                                'region': zone['region'], 'region_name': zone['name'],
                                'period_start': start.isoformat(), 'period_end': end.isoformat(), 'period_end_exclusive': True,
                                'source': 'EOG', 'product_declared': 'monthly/v10/vcmslcfg',
                                'satellite_declared': 'J02' if month == '202511' else 'NPP',
                                'satellite_verified': False, 'unit_declared': 'nW/cm²/sr',
                                'statistics': stats, 'flags': review_flags(month, stats),
                                'boundary_sha256': boundary_hashes, 'all_touched': False,
                                'method': 'raw signed radiance; cf>0 integer; no interpolation, clipping, smoothing or calibration'})
            del rad, cf
        if sha256(path) != digest or (bbox and sha256(bbox) != bbox_digest):
            raise ValueError('影像在检查时变化，拒绝提交本期')
        if any(sha256(p) != boundary_hashes[p.name] for p in boundary_paths):
            raise ValueError('边界在检查时变化')
        asset = asset_for(root, path, digest, 'raster')
        report = {'asset': asset, 'grid': grid, 'bbox_sha256': bbox_digest,
                  'bbox_path': bbox.relative_to(root).as_posix() if bbox else None,
                  'inside_equal_bbox': equal, 'outside_nodata_correct': outside_correct,
                  'boundaries': boundaries, 'records': records}
        report_path = export_json(output / 'rasters', fingerprint(report) + '.json', report)
        check = {'kind': 'raster_check', 'location': month, 'status': 'pending_review',
                 **{k: report[k] for k in ('grid', 'bbox_sha256', 'bbox_path', 'inside_equal_bbox', 'outside_nodata_correct', 'boundaries')}}
        added = store.save(asset, [check, *records]) if store else 0
        results.append({'month': month, 'report': report_path, 'national': records[0]['statistics'],
                        'flags': records[0]['flags'], 'equal_bbox': equal, 'outside_correct': outside_correct,
                        'new_records': added})
        print(f'影像 {index}/{len(main)} {month}: 界内逐像元一致={equal}', flush=True)
    expected = {f'{y}{m:02}' for y in range(2020, 2027) for m in range(1, 13) if f'{y}{m:02}' <= '202608'}
    gaps = sorted(expected - present)
    if present - expected:
        raise ValueError('有超出本次交付范围的月份')
    # 缺月只记清单，不创建数值为零的正式观测。
    gap_asset = {'path': 'expected_month_manifest', 'kind': 'manifest', 'sha256': fingerprint(sorted(present)),
                 'review_version': VERSION}
    gap_asset['id'] = fingerprint(gap_asset)
    if store:
        store.save(gap_asset, [{'kind': 'missing_month', 'location': m, 'status': 'missing',
                               'value': None, 'reason': 'not delivered; upstream cause not independently verified'} for m in gaps])
    return {'months': results, 'missing_months': gaps, 'boundaries': boundaries,
            'all_grid_equal': True, 'originals_unchanged': True}


def commit_raster_audit(root, output, audit_path, store):
    """提交已经完整复算的报告，逐一重验报告摘要与输入文件摘要，不重复计算像元。"""
    audit_path = audit_path.resolve(strict=True)
    if not audit_path.is_relative_to(output):
        raise ValueError('只接受本待审目录的核验汇总')
    audit = json.loads(audit_path.read_text(encoding='utf-8'))
    if audit_path.name != 'rasters-' + fingerprint(audit)[:16] + '.json' or audit['version'] != VERSION:
        raise ValueError('核验汇总摘要或版本无效')
    reports = []
    for item in audit['result']['months']:
        path = Path(item['report']).resolve(strict=True)
        if not path.is_relative_to(output / 'rasters'):
            raise ValueError('核验报告越界')
        report = json.loads(path.read_text(encoding='utf-8'))
        if path.stem != fingerprint(report):
            raise ValueError('核验报告摘要不匹配')
        asset = report['asset']
        if fingerprint({k: v for k, v in asset.items() if k != 'id'}) != asset['id']:
            raise ValueError('资产摘要不匹配')
        source = (root / asset['path']).resolve(strict=True)
        bbox = (root / report['bbox_path']).resolve(strict=True) if report['bbox_path'] else None
        if not source.is_relative_to(root) or (bbox and not bbox.is_relative_to(root)):
            raise ValueError('源文件越界')
        if sha256(source) != asset['sha256'] or (bbox and sha256(bbox) != report['bbox_sha256']):
            raise ValueError('原件已变化，必须重新核验')
        for name, digest in report['boundaries']['hashes'].items():
            if Path(name).name != name or sha256(source.parent.parent / 'boundary' / name) != digest:
                raise ValueError('边界版本已变化')
        reports.append(report)
    inserted = 0
    for report in reports:
        records = report['records']
        check = {'kind': 'raster_check', 'location': records[0]['period_start'][:7].replace('-', ''), 'status': 'pending_review',
                 **{k: report[k] for k in ('grid', 'bbox_sha256', 'bbox_path', 'inside_equal_bbox', 'outside_nodata_correct', 'boundaries')}}
        inserted += store.save(report['asset'], [check, *records])
    present = sorted(r['records'][0]['period_start'][:7].replace('-', '') for r in reports)
    gap_asset = {'path': 'expected_month_manifest', 'kind': 'manifest', 'sha256': fingerprint(present), 'review_version': VERSION}
    gap_asset['id'] = fingerprint(gap_asset)
    inserted += store.save(gap_asset, [{'kind': 'missing_month', 'location': m, 'status': 'missing', 'value': None,
                                      'reason': 'not delivered; upstream cause not independently verified'} for m in audit['result']['missing_months']])
    return {'raster_files_verified': len(reports), 'new_records': inserted, 'originals_hash_rechecked': True}


def verify_store(root, output):
    """只读数据库回读：数量、校验摘要、原件、候选质量与表格口径。"""
    connection = sqlite3.connect((output / 'review.sqlite3').as_uri() + '?mode=ro', uri=True)
    try:
        if connection.execute('PRAGMA integrity_check').fetchone()[0] != 'ok' or connection.execute('PRAGMA foreign_key_check').fetchall():
            raise ValueError('待审库完整性检查失败')
        assets = [json.loads(r[0]) for r in connection.execute('SELECT payload FROM assets')]
        rows = [json.loads(r[0]) for r in connection.execute('SELECT payload FROM pending_records')]
        for asset in assets:
            if asset['kind'] != 'manifest':
                path = (root / asset['path']).resolve(strict=True)
                if not path.is_relative_to(root) or sha256(path) != asset['sha256']:
                    raise ValueError('入库后原件摘要变化')
        if connection.execute('SELECT count(*) FROM pending_records WHERE eligible_for_scoring!=0').fetchone()[0] or any(row['eligible_for_scoring'] is not False for row in rows):
            raise ValueError('待审隔离标记失效')
        stats = [r for r in rows if r['kind'] == 'nightlight_candidate']
        national = [r for r in stats if r['region'] == 'MMR']
        raster_checks = [r for r in rows if r['kind'] == 'raster_check']
        cases = {r['period_start']: r['statistics'] for r in national if r['period_start'] in
                 {'2020-01-01', '2022-07-01', '2025-02-01', '2025-11-01'}}
        tables = []
        for asset in assets:
            if not asset['path'].endswith('.xlsx'):
                continue
            blocks = [json.loads(r[0]) for r in connection.execute('SELECT payload FROM pending_records WHERE asset_id=?', (asset['id'],))]
            sheets = {}
            for block in blocks:
                sheet, number = block['location'].rsplit(':row:', 1)
                sheets.setdefault(sheet, []).append((int(number), {re.sub(r'\d', '', c['cell']): c['value'] for c in json.loads(block['text'])}))
            info = {'file': asset['path']}
            if '文献总表' in sheets:
                ordered = sorted(sheets['文献总表'])
                headers = ordered[0][1]
                body = [{headers.get(k, k): v for k, v in row.items()} for _, row in ordered[1:]]
                body = [r for r in body if r.get('题名')]
                info.update(catalogue_rows=len(body),
                            missing_fields={k: sum(not r.get(k) for r in body) for k in ('作者', '年份', '摘要', 'DOI')},
                            declared_review_status=dict(Counter(str(r.get('核查状态', '未提供')) for r in body)))
            if '①标注表' in sheets:
                filled = [(n, r) for n, r in sheets['①标注表'] if n > 1 and isinstance(r.get('A'), (int, float)) and r.get('B')]
                info.update(annotation_rows=len(filled), blank_annotators=sum(not r.get('P') for _, r in filled),
                            no_source_url=sum(not str(r.get('O', '')).startswith(('http://', 'https://')) for _, r in filled),
                            uncertain_date_rows=[n for n, r in filled if '待考' in str(r.get('Q', '')) or '按当月' in str(r.get('Q', ''))])
            if '权重分配表' in sheets:
                weights = {r['B'].strip(): r['D'] for _, r in sheets['权重分配表'] if r.get('A') == '指标层'}
                info.update(global_weights=weights, global_weight_sum=round(sum(weights.values()), 12), approved=False)
            tables.append(info)
        return {'integrity_check': 'ok', 'source_assets_hash_verified': sum(a['kind'] != 'manifest' for a in assets),
                'asset_count': len(assets), 'record_counts': dict(Counter(r['kind'] for r in rows)),
                'eligible_for_scoring': 0, 'regions': sorted({r['region'] for r in stats}),
                'raster_month_count': len(national), 'all_bbox_equal': all(r['inside_equal_bbox'] for r in raster_checks) if raster_checks else None,
                'all_outside_nodata': all(r['outside_nodata_correct'] for r in raster_checks) if raster_checks else None,
                'candidate_status_counts': dict(Counter(r['status'] for r in stats)),
                'missing_months': [r['location'] for r in rows if r['kind'] == 'missing_month'],
                'low_coverage_national_months': [r['period_start'] for r in national if (r['statistics']['observed_area_fraction'] or 0) < .5],
                'national_case_statistics': cases,
                'boundary_audit': raster_checks[0]['boundaries'] if raster_checks else None,
                'tables': tables, 'image_ocr_completed': False, 'literature_fulltexts_verified': False}
    finally:
        connection.close()


def formal_output_hashes(root):
    """仅核对正式文件，既不初始化系统，也不连接共享数据库。"""
    runtime = root / '运行数据'
    names = ['processed/daily_risk.jsonl', 'processed/risk_scores.jsonl',
             'external/risk_state.json', 'external/indicator_observations.jsonl',
             'processed/gdelt_event_store.json']
    return {name: sha256(runtime / name) if (runtime / name).is_file() else None for name in names}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source-root', type=Path, required=True)
    parser.add_argument('--output-root', type=Path, required=True)
    parser.add_argument('--phase', choices=('documents', 'rasters', 'commit-rasters', 'verify'), required=True)
    parser.add_argument('--audit-summary', type=Path)
    parser.add_argument('--apply', action='store_true')
    args = parser.parse_args()
    root = args.source_root.resolve(strict=True)
    output = args.output_root.resolve()
    # 只写独立审核子目录，不允许把报告放进交付的影像/正文目录。
    if output == root or output.name != 'team_delivery_review':
        parser.error('output-root必须为独立的team_delivery_review目录')
    protected_before = formal_output_hashes(root)
    store = ReviewStore(output) if args.apply else None
    try:
        if args.phase == 'verify':
            result = verify_store(root, output)
        elif args.phase == 'commit-rasters':
            if not args.apply or not args.audit_summary:
                parser.error('提交已审影像需要--apply与--audit-summary')
            result = commit_raster_audit(root, output, args.audit_summary, store)
        else:
            result = process_documents(root, output, store) if args.phase == 'documents' else process_rasters(root, output, store)
        protected_after = formal_output_hashes(root)
        summary = {'phase': args.phase, 'applied': args.apply, 'version': VERSION,
                   'result': result, 'counts': store.counts() if store else None,
                   'formal_outputs_unchanged': protected_before == protected_after,
                   'formal_hashes_before': protected_before, 'formal_hashes_after': protected_after}
        path = export_json(output, args.phase + '-' + fingerprint(summary)[:16] + '.json', summary)
        print('汇总：' + path)
        print('待审库：' + encoded(summary['counts']))
    finally:
        if store:
            store.close()


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    main()
